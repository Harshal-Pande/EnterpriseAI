"""Seed load and database validation integration tests."""

from pathlib import Path

import pytest

from column_mappings import EXPECTED_SOURCE_ROW_COUNTS, SHEET_MAPPINGS  # noqa: E402
from load_seed_data import load_workbook, sheet_to_records  # noqa: E402
from migrate import run  # noqa: E402
from validate_seed_data import validate  # noqa: E402


def _prepare(monkeypatch, test_database_url):
    monkeypatch.setenv("TEST_DATABASE_URL", test_database_url)
    assert run(dry_run=False, env_var="TEST_DATABASE_URL") == 0
    from db.connection import connect
    c = connect("TEST_DATABASE_URL")
    with c.cursor() as cur:
        cur.execute("SET app.allow_event_log_mutation = true;")
        cur.execute("DELETE FROM evidence_references WHERE evidence_id IS NOT NULL; DELETE FROM agent_tool_calls WHERE tool_call_id IS NOT NULL; DELETE FROM inter_agent_messages WHERE message_id IS NOT NULL; DELETE FROM task_queue WHERE task_id IS NOT NULL; DELETE FROM task_batches WHERE batch_id IS NOT NULL; DELETE FROM event_log WHERE event_id IS NOT NULL; DELETE FROM action_queue WHERE action_id IS NOT NULL; DELETE FROM policy_evaluations WHERE policy_eval_id IS NOT NULL; DELETE FROM approval_requests WHERE approval_id IS NOT NULL; DELETE FROM agent_runs WHERE agent_run_id IS NOT NULL; DELETE FROM decisions WHERE decision_id IS NOT NULL; DELETE FROM scenario_instances WHERE scenario_id IS NOT NULL;")
        c.commit()
    c.close()


@pytest.mark.integration
def test_seed_row_counts_and_dq01(test_conn, test_database_url, monkeypatch, workbook_path):
    import openpyxl

    _prepare(monkeypatch, test_database_url)
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        result = load_workbook(wb, test_conn, dry_run=False)
    finally:
        wb.close()
    assert result["totals"]["rejected"] == 0
    assert result["totals"]["source"] == 1327

    with test_conn.cursor() as cur:
        for sheet, mapping in SHEET_MAPPINGS.items():
            cur.execute(f"SELECT COUNT(*) FROM {mapping.dest_table}")
            assert cur.fetchone()[0] == EXPECTED_SOURCE_ROW_COUNTS[sheet]
        cur.execute(
            "SELECT COUNT(*) FROM scenario_templates WHERE linked_reference_id IS NULL AND enabled = TRUE"
        )
        assert cur.fetchone()[0] == 15
        cur.execute("SELECT COUNT(*) FROM scenario_instances")
        assert cur.fetchone()[0] == 0


@pytest.mark.integration
def test_seed_is_idempotent(test_conn, test_database_url, monkeypatch, workbook_path):
    import openpyxl

    _prepare(monkeypatch, test_database_url)
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        first = load_workbook(wb, test_conn, dry_run=False)
        wb.close()
        wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        second = load_workbook(wb, test_conn, dry_run=False)
    finally:
        wb.close()

    assert first["totals"]["rejected"] == 0
    assert second["totals"]["inserted"] == 0
    assert second["totals"]["skipped"] == 1327
    assert second["totals"]["updated"] == 0


@pytest.mark.integration
def test_database_validation_report(test_conn, test_database_url, monkeypatch, workbook_path):
    import openpyxl

    _prepare(monkeypatch, test_database_url)
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        load_workbook(wb, test_conn, dry_run=False)
    finally:
        wb.close()
    report = validate(test_conn)
    assert report["database_queried"] is True
    assert report["pass"] is True


@pytest.mark.integration
def test_duplicate_primary_key_is_skipped_not_duplicated(
    test_conn, test_database_url, monkeypatch, workbook_path
):
    import openpyxl

    _prepare(monkeypatch, test_database_url)
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        load_workbook(wb, test_conn, dry_run=False)
    finally:
        wb.close()
    with test_conn.cursor() as cur:
        cur.execute(
            "INSERT INTO entities (entity_id, entity_name, entity_type, status) "
            "VALUES ('EAS-HQ', 'duplicate attempt', 'Corporate', 'ACTIVE') "
            "ON CONFLICT (entity_id) DO NOTHING"
        )
        assert cur.rowcount == 0
        cur.execute("SELECT entity_name FROM entities WHERE entity_id = 'EAS-HQ'")
        name = cur.fetchone()[0]
        assert name != "duplicate attempt"
    test_conn.commit()


@pytest.mark.integration
def test_invalid_record_rolls_back_in_transaction(test_conn, test_database_url, monkeypatch):
    import psycopg2

    _prepare(monkeypatch, test_database_url)
    with test_conn.cursor() as cur:
        cur.execute("SELECT COUNT(*) FROM products")
        before = cur.fetchone()[0]
        try:
            cur.execute(
                "INSERT INTO products (product_id, product_name, product_category, base_uom, status) "
                "VALUES ('P-TMP', 'Tmp', 'Flux', 'MT', 'ACTIVE')"
            )
            cur.execute(
                "INSERT INTO products (product_id, product_name, product_category, base_uom, status) "
                "VALUES ('P-TMP-2', 'Tmp2', 'Flux', 'MT', 'NOT_A_STATUS')"
            )
            test_conn.commit()
            raise AssertionError("expected check violation")
        except psycopg2.errors.CheckViolation:
            test_conn.rollback()
        cur.execute("SELECT COUNT(*) FROM products")
        after = cur.fetchone()[0]
        assert after == before


def test_workbook_sheet_helpers_read_source_rows(workbook_path):
    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        entities = sheet_to_records(wb["ref_entities"])
        empty = sheet_to_records(wb["scenario_instances"])
    finally:
        wb.close()
    assert len(entities) == 4
    assert empty == []
    assert Path(workbook_path).name.endswith(".xlsx")
