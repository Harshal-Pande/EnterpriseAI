"""Mapping completeness and DQ-01 transformation tests (no PostgreSQL)."""

from column_mappings import (
    EMPTY_RUNTIME_SHEETS,
    EXCLUDED_SHEETS,
    EXPECTED_SOURCE_ROW_COUNTS,
    FK_SPECS,
    LOAD_SEQUENCE,
    MAPPED_IMPORT_ROWS,
    METADATA_EXCLUDED_ROWS,
    SHEET_MAPPINGS,
    TOTAL_WORKBOOK_ROWS,
    _severity_from_linked_ref,
    convert_record,
    dest_primary_key,
)


def test_count_reconciliation_constants():
    assert TOTAL_WORKBOOK_ROWS == 1677
    assert MAPPED_IMPORT_ROWS == 1327
    assert METADATA_EXCLUDED_ROWS == 350
    assert 1677 == 1327 + 350 + 0
    assert len(SHEET_MAPPINGS) == 31
    assert len(EMPTY_RUNTIME_SHEETS) == 6
    assert sum(EXPECTED_SOURCE_ROW_COUNTS.values()) == 1327


def test_load_sequence_covers_all_mapped_sheets():
    assert set(LOAD_SEQUENCE) == set(SHEET_MAPPINGS)
    assert LOAD_SEQUENCE[0] == "ref_entities"
    assert LOAD_SEQUENCE.index("ref_products") < LOAD_SEQUENCE.index("inventory")
    assert LOAD_SEQUENCE.index("orders") < LOAD_SEQUENCE.index("order_lines")
    assert LOAD_SEQUENCE.index("shipments") < LOAD_SEQUENCE.index("carrier_events")


def test_expected_counts_match_mappings():
    assert set(EXPECTED_SOURCE_ROW_COUNTS) == set(SHEET_MAPPINGS)


def test_dq01_severity_transform():
    assert _severity_from_linked_ref("HIGH") == "HIGH"
    assert _severity_from_linked_ref("not-a-severity") is None
    mapping = SHEET_MAPPINGS["scenario_templates"]
    dest, err = convert_record(
        mapping,
        {
            "scenario_template_id": "TPL-001",
            "scenario_type": "INVENTORY_SHORTAGE",
            "domain": "Inventory",
            "description": "demo",
            "linked_reference_id": "HIGH",
            "base_record_id": "INV-1",
            "injection_method": "overlay",
            "agents_involved": "inventory",
            "expected_response": "reorder",
            "guardrail": "no auto buy",
        },
    )
    assert err is None
    assert dest["severity"] == "HIGH"
    assert dest["linked_reference_id"] is None
    assert dest["enabled"] is True


def test_convert_record_rejects_missing_required():
    mapping = SHEET_MAPPINGS["ref_entities"]
    dest, err = convert_record(mapping, {"entity_name": "x"})
    assert dest is None
    assert "entity_id" in err


def test_blank_string_becomes_null_on_optional():
    mapping = SHEET_MAPPINGS["ref_entities"]
    dest, err = convert_record(
        mapping,
        {
            "entity_id": "EAS-HQ",
            "entity_name": "HQ",
            "entity_type": "Corporate",
            "country": "  ",
            "status": "ACTIVE",
        },
    )
    assert err is None
    assert dest["country"] is None


def test_fk_specs_point_at_mapped_sheets():
    for spec in FK_SPECS:
        assert spec.child_sheet in SHEET_MAPPINGS
        assert spec.parent_sheet in SHEET_MAPPINGS


def test_excluded_and_runtime_sheets_are_not_imported():
    for name in EXCLUDED_SHEETS:
        assert name not in SHEET_MAPPINGS
    for name in EMPTY_RUNTIME_SHEETS:
        assert name in EXCLUDED_SHEETS


def test_source_columns_exist_in_workbook(workbook_path):
    openpyxl = pytest_import_openpyxl()
    if not workbook_path.exists():
        raise AssertionError(f"Workbook missing: {workbook_path}")
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        for sheet_name, mapping in SHEET_MAPPINGS.items():
            assert sheet_name in wb.sheetnames
            rows = list(wb[sheet_name].iter_rows(values_only=True))
            headers = [str(h) for h in rows[0]]
            for col in mapping.columns:
                assert col.source_col in headers, (
                    f"{sheet_name} missing mapped column {col.source_col}"
                )
            assert dest_primary_key(mapping)
        for name in EMPTY_RUNTIME_SHEETS:
            rows = list(wb[name].iter_rows(values_only=True))
            assert len(rows) <= 1 or all(
                all(v is None for v in row) for row in rows[1:]
            )
            data_rows = [r for r in rows[1:] if any(v is not None for v in r)]
            assert data_rows == []
    finally:
        wb.close()


def pytest_import_openpyxl():
    import openpyxl

    return openpyxl
