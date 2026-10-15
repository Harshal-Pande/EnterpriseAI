"""Load the industrial workbook into PostgreSQL using column_mappings.py.

Usage:
    python backend/data/load_seed_data.py --dry-run
    python backend/data/load_seed_data.py
    python backend/data/load_seed_data.py --database-url-var TEST_DATABASE_URL

Idempotent: INSERT ... ON CONFLICT (pk) DO NOTHING. Existing operational
rows are never overwritten.
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install -r backend/requirements.txt")
    sys.exit(1)

_HERE = Path(__file__).resolve().parent
_BACKEND = _HERE.parent
for path in (_HERE, _BACKEND, _BACKEND / "db"):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from column_mappings import (  # noqa: E402
    EXPECTED_SOURCE_ROW_COUNTS,
    FK_SPECS,
    LOAD_SEQUENCE,
    SHEET_MAPPINGS,
    convert_record,
    dest_primary_key,
)
from connection import connect, get_database_url  # noqa: E402

def _find_workbook() -> Path:
    candidates = [
        Path(__file__).resolve().parent / "enterprise_ai_industrial_dataset_v2.xlsx",
        Path(__file__).resolve().parent.parent.parent / "project_reference" / "enterprise_ai_industrial_dataset_v2.xlsx",
        Path(__file__).resolve().parent.parent / "project_reference" / "enterprise_ai_industrial_dataset_v2.xlsx",
        Path("/app/data/enterprise_ai_industrial_dataset_v2.xlsx"),
        Path("/app/project_reference/enterprise_ai_industrial_dataset_v2.xlsx"),
    ]
    for p in candidates:
        if p.is_file():
            return p
    return candidates[0]


WORKBOOK_PATH = _find_workbook()



def sheet_to_records(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) if h is not None else f"COL_{i}" for i, h in enumerate(rows[0])]
    records = []
    for row in rows[1:]:
        rec = {}
        for i, val in enumerate(row):
            col = headers[i] if i < len(headers) else f"COL_{i}"
            rec[col] = val
        records.append(rec)
    return records


def _parent_table_and_pk(parent_sheet: str) -> tuple[str, str]:
    mapping = SHEET_MAPPINGS[parent_sheet]
    return mapping.dest_table, dest_primary_key(mapping)


def _child_dest_col(child_sheet: str, source_fk: str) -> str:
    mapping = SHEET_MAPPINGS[child_sheet]
    for col in mapping.columns:
        if col.source_col == source_fk:
            return col.dest_col
    return source_fk


def insert_sql(table: str, columns: list[str], pk: str) -> str:
    cols = ", ".join(columns)
    placeholders = ", ".join(["%s"] * len(columns))
    return (
        f"INSERT INTO {table} ({cols}) VALUES ({placeholders}) "
        f"ON CONFLICT ({pk}) DO NOTHING"
    )


def load_workbook(wb, conn, dry_run: bool) -> dict[str, Any]:
    pk_sets: dict[str, set[str]] = defaultdict(set)
    summaries = []
    totals = {"source": 0, "inserted": 0, "updated": 0, "skipped": 0, "rejected": 0}

    # Preload existing PKs so FK checks work on reruns.
    if conn is not None:
        with conn.cursor() as cur:
            for sheet_name in LOAD_SEQUENCE:
                mapping = SHEET_MAPPINGS[sheet_name]
                table = mapping.dest_table
                pk = dest_primary_key(mapping)
                cur.execute(f"SELECT {pk} FROM {table}")
                pk_sets[table] = {str(row[0]) for row in cur.fetchall()}

    for sheet_name in LOAD_SEQUENCE:
        mapping = SHEET_MAPPINGS[sheet_name]
        records = sheet_to_records(wb[sheet_name])
        table = mapping.dest_table
        pk_col = dest_primary_key(mapping)
        dest_columns = [c.dest_col for c in mapping.columns]
        for extra in mapping.extra_dest_values:
            if extra not in dest_columns:
                dest_columns.append(extra)

        inserted = skipped = rejected = updated = 0
        rejected_examples: list[str] = []

        converted_rows: list[dict] = []
        for rec in records:
            dest, reason = convert_record(mapping, rec)
            if dest is None:
                rejected += 1
                if len(rejected_examples) < 5:
                    rejected_examples.append(reason or "conversion failed")
                continue

            fk_fail = None
            for spec in FK_SPECS:
                if spec.child_sheet != sheet_name:
                    continue
                child_col = _child_dest_col(sheet_name, spec.fk_col)
                value = dest.get(child_col)
                if value is None:
                    continue
                parent_table, parent_pk = _parent_table_and_pk(spec.parent_sheet)
                if str(value) not in pk_sets[parent_table]:
                    fk_fail = (
                        f"FK {child_col}={value} not in {parent_table}.{parent_pk}"
                    )
                    break
            if fk_fail:
                rejected += 1
                if len(rejected_examples) < 5:
                    rejected_examples.append(fk_fail)
                continue
            converted_rows.append(dest)

        if dry_run or conn is None:
            for dest in converted_rows:
                key = str(dest[pk_col])
                if key in pk_sets[table]:
                    skipped += 1
                else:
                    inserted += 1
                    pk_sets[table].add(key)
        else:
            sql = insert_sql(table, dest_columns, pk_col)
            with conn.cursor() as cur:
                for dest in converted_rows:
                    key = str(dest[pk_col])
                    values = [dest.get(col) for col in dest_columns]
                    cur.execute(sql, values)
                    if cur.rowcount == 1:
                        inserted += 1
                        pk_sets[table].add(key)
                    else:
                        skipped += 1
                        pk_sets[table].add(key)

        expected = EXPECTED_SOURCE_ROW_COUNTS.get(sheet_name)
        summary = {
            "sheet": sheet_name,
            "table": table,
            "source_rows": len(records),
            "expected_source_rows": expected,
            "inserted": inserted,
            "updated": updated,
            "skipped": skipped,
            "rejected": rejected,
            "rejected_examples": rejected_examples,
        }
        summaries.append(summary)
        totals["source"] += len(records)
        totals["inserted"] += inserted
        totals["updated"] += updated
        totals["skipped"] += skipped
        totals["rejected"] += rejected

        print(
            f"{sheet_name:28s} -> {table:24s} "
            f"source={len(records):4d} inserted={inserted:4d} "
            f"skipped={skipped:4d} rejected={rejected:4d} updated={updated}"
        )
        if rejected_examples:
            for ex in rejected_examples:
                print(f"    reject: {ex}")

    if conn is not None and not dry_run:
        if totals["rejected"] > 0:
            conn.rollback()
        else:
            conn.commit()

    return {"tables": summaries, "totals": totals, "dry_run": dry_run}


def main() -> None:
    parser = argparse.ArgumentParser(description="Load workbook seed data into PostgreSQL.")
    parser.add_argument("--dry-run", action="store_true", help="Validate and count without writing.")
    parser.add_argument("--database-url-var", default="DATABASE_URL")
    args = parser.parse_args()

    if not WORKBOOK_PATH.exists():
        print(f"ERROR: Workbook not found at {WORKBOOK_PATH}", file=sys.stderr)
        sys.exit(1)

    print(f"Workbook: {WORKBOOK_PATH}")
    wb = openpyxl.load_workbook(str(WORKBOOK_PATH), read_only=True, data_only=True)

    conn = None
    if args.dry_run:
        url = get_database_url(args.database_url_var)
        if url:
            conn = connect(args.database_url_var)
            print("DRY-RUN with database (existing PKs loaded; no writes).")
        else:
            print("DRY-RUN without database (FK checks use in-memory inserted keys only).")
    else:
        conn = connect(args.database_url_var)

    try:
        result = load_workbook(wb, conn, dry_run=args.dry_run)
    except Exception:
        if conn is not None:
            conn.rollback()
        raise
    finally:
        wb.close()
        if conn is not None:
            conn.close()

    totals = result["totals"]
    print("\nTotals")
    print(f"  source   : {totals['source']}")
    print(f"  inserted : {totals['inserted']}")
    print(f"  skipped  : {totals['skipped']}")
    print(f"  rejected : {totals['rejected']}")
    print(f"  updated  : {totals['updated']} (Phase 1 never overwrites existing rows)")
    if totals["rejected"] > 0:
        sys.exit(1)


if __name__ == "__main__":
    main()
