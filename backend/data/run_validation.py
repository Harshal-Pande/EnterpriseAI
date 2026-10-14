"""
run_validation.py — Workbook-only validation (no PostgreSQL required).

Reads every mapped sheet from the workbook, applies validator rules, and writes
backend/data/workbook_validation_report.json.

Usage:
    python backend/data/run_validation.py
"""

from __future__ import annotations

import json
import sys
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed.  Run: pip install openpyxl==3.1.5")
    sys.exit(1)

_HERE = Path(__file__).resolve().parent
if str(_HERE) not in sys.path:
    sys.path.insert(0, str(_HERE))

from column_mappings import (  # noqa: E402
    EMPTY_RUNTIME_SHEETS,
    EXPECTED_SOURCE_ROW_COUNTS,
    EXCLUDED_SHEETS,
    FK_SPECS,
    MAPPED_IMPORT_ROWS,
    METADATA_EXCLUDED_ROWS,
    SHEET_MAPPINGS,
    STATUS_ALLOWLISTS,
    TOTAL_WORKBOOK_ROWS,
)
from validators import (  # noqa: E402
    ValidationError,
    check_arithmetic_identity,
    check_date_order,
    check_date_range,
    check_no_duplicate_ids,
    check_no_nulls,
    check_numeric_range,
    check_product_identity,
    check_quantity_lte,
    check_referential_integrity,
    check_required_columns,
    check_valid_status,
    _to_float,
)

WORKBOOK_PATH = (
    Path(__file__).resolve().parent.parent.parent
    / "project_reference"
    / "enterprise_ai_industrial_dataset_v2.xlsx"
)
OUTPUT_PATH = Path(__file__).resolve().parent / "workbook_validation_report.json"

MIN_DATE = date(2025, 1, 1)
MAX_DATE = date(2028, 12, 31)

NUMERIC_RANGES = {
    ("ref_products", "reference_unit_cost_inr"): (0, 10_000_000),
    ("ref_warehouses", "capacity_units"): (0, 1_000_000),
    ("ref_suppliers", "historical_on_time_rate"): (0.0, 1.0),
    ("ref_suppliers", "default_lead_time_hours"): (0, 8760),
    ("ref_supplier_offers", "unit_price_inr"): (0, 10_000_000),
    ("ref_supplier_offers", "on_time_rate"): (0.0, 1.0),
    ("ref_production_lines", "utilisation_pct"): (0.0, 1.0),
    ("ref_production_lines", "daily_capacity_mt"): (0, 50_000),
    ("ref_budgets", "approved_limit_inr"): (0, 10_000_000_000),
    ("ref_budgets", "committed_spend_inr"): (0, 10_000_000_000),
    ("ref_budgets", "available_budget_inr"): (0, 10_000_000_000),
    ("ref_shifts", "duration_hours"): (1, 24),
    ("ref_reorder_policies", "reorder_point_qty"): (0, 1_000_000),
    ("ref_reorder_policies", "safety_stock_qty"): (0, 1_000_000),
    ("bill_of_materials", "quantity_per_output_unit"): (0.0001, 1000),
    ("inventory", "on_hand_qty"): (0, 10_000_000),
    ("inventory", "reserved_qty"): (0, 10_000_000),
    ("inventory", "available_qty"): (0, 10_000_000),
    ("order_lines", "requested_qty"): (0.001, 1_000_000),
    ("order_lines", "unit_price_inr"): (0, 100_000_000),
    ("production_batches", "planned_output_qty_mt"): (0, 100_000),
    ("production_batches", "actual_output_qty_mt"): (0, 100_000),
    ("purchase_orders", "quantity"): (0.001, 10_000_000),
    ("purchase_orders", "unit_cost_inr"): (0, 10_000_000),
    ("goods_receipts", "received_qty"): (0, 1_000_000),
    ("goods_receipts", "accepted_qty"): (0, 1_000_000),
    ("maintenance_assets", "failure_risk_score"): (0.0, 1.0),
    ("energy_readings", "energy_kwh"): (0, 1_000_000),
    ("energy_readings", "unit_rate_inr"): (0, 100),
    ("feed_health", "consecutive_failures"): (0, 1000),
}

DATE_COLUMNS = {
    ("ref_supplier_offers", "effective_from"),
    ("ref_supplier_offers", "effective_to"),
    ("ref_supplier_contracts", "effective_from"),
    ("ref_supplier_contracts", "effective_to"),
    ("orders", "created_at"),
    ("orders", "required_delivery_at"),
    ("production_batches", "planned_start"),
    ("production_batches", "planned_end"),
    ("shipments", "planned_dispatch_at"),
    ("shipments", "estimated_arrival_at"),
    ("quality_events", "detected_at"),
    ("maintenance_assets", "next_maintenance_due"),
    ("maintenance_work_orders", "scheduled_for"),
    ("energy_readings", "reading_at"),
    ("safety_events", "occurred_at"),
    ("carrier_events", "event_timestamp"),
}


def sheet_to_records(ws) -> tuple[list[str], list[dict]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h) if h is not None else f"COL_{i}" for i, h in enumerate(rows[0])]
    records = []
    for row in rows[1:]:
        rec = {}
        for i, val in enumerate(row):
            col = headers[i] if i < len(headers) else f"COL_{i}"
            rec[col] = val
        records.append(rec)
    return headers, records


def build_ref_set(records: list[dict], col: str) -> set:
    return {str(r[col]).strip() for r in records if r.get(col) is not None and str(r[col]).strip() != ""}


def errors_to_findings(errors: list[ValidationError]) -> list[dict]:
    grouped: dict = defaultdict(
        lambda: {"rule": "", "severity": "", "column": "", "count": 0, "examples": []}
    )
    for e in errors:
        key = (e.rule, e.column, e.severity)
        g = grouped[key]
        g["rule"] = e.rule
        g["severity"] = e.severity
        g["column"] = e.column
        g["count"] += 1
        if len(g["examples"]) < 5:
            g["examples"].append(str(e.value))
    return list(grouped.values())


def _info(sheet: str, column: str, rule: str, message: str, value=None, row: int = 0) -> ValidationError:
    return ValidationError(
        sheet=sheet, row=row, column=column, value=value,
        rule=rule, severity="INFO", message=message,
    )


def validate_all(wb) -> dict:
    all_errors: list[ValidationError] = []
    sheet_results: list[dict] = []
    loaded: dict[str, list[dict]] = {}
    headers_by_sheet: dict[str, list[str]] = {}

    for name in list(SHEET_MAPPINGS) + list(EMPTY_RUNTIME_SHEETS) + list(EXCLUDED_SHEETS):
        if name not in wb.sheetnames:
            continue
        cols, recs = sheet_to_records(wb[name])
        loaded[name] = recs
        headers_by_sheet[name] = cols

    parent_sets = {}
    for spec in FK_SPECS:
        parent_sets[(spec.parent_sheet, spec.parent_pk)] = build_ref_set(
            loaded.get(spec.parent_sheet, []), spec.parent_pk
        )

    def run_sheet(sheet_name: str, errs: list[ValidationError]):
        recs = loaded.get(sheet_name, [])
        cols = headers_by_sheet.get(sheet_name, [])
        findings = errors_to_findings(errs)
        all_errors.extend(errs)
        sheet_results.append({
            "sheet_name": sheet_name,
            "row_count": len(recs),
            "column_count": len(cols),
            "columns": cols,
            "findings": findings,
        })

    for sheet_name, mapping in SHEET_MAPPINGS.items():
        recs = loaded.get(sheet_name, [])
        cols = headers_by_sheet.get(sheet_name, [])
        errs: list[ValidationError] = []

        required_src = [c.source_col for c in mapping.columns if c.required]
        errs += check_required_columns(sheet_name, cols, required_src)
        errs += check_no_duplicate_ids(sheet_name, recs, mapping.primary_key_col)
        errs += check_no_nulls(sheet_name, recs, required_src)

        for (s, col), allowed in STATUS_ALLOWLISTS.items():
            if s == sheet_name:
                errs += check_valid_status(sheet_name, recs, col, allowed)

        for (s, col), (lo, hi) in NUMERIC_RANGES.items():
            if s == sheet_name:
                errs += check_numeric_range(sheet_name, recs, col, lo, hi)

        for s, col in DATE_COLUMNS:
            if s == sheet_name:
                errs += check_date_range(sheet_name, recs, col, MIN_DATE, MAX_DATE)

        for spec in FK_SPECS:
            if spec.child_sheet == sheet_name:
                ref = parent_sets.get((spec.parent_sheet, spec.parent_pk), set())
                errs += check_referential_integrity(sheet_name, recs, spec.fk_col, ref)

        if sheet_name == "inventory":
            errs += check_arithmetic_identity(
                sheet_name, recs, "on_hand_qty", "reserved_qty", "available_qty"
            )
        if sheet_name == "ref_budgets":
            errs += check_arithmetic_identity(
                sheet_name, recs, "approved_limit_inr", "committed_spend_inr", "available_budget_inr"
            )
        if sheet_name == "purchase_orders":
            errs += check_product_identity(
                sheet_name, recs, "quantity", "unit_cost_inr", "total_cost_inr"
            )
        if sheet_name == "goods_receipts":
            errs += check_quantity_lte(sheet_name, recs, "accepted_qty", "received_qty")
            for idx, rec in enumerate(recs, start=1):
                if str(rec.get("grn_id")) == "GRN-002":
                    errs.append(_info(
                        sheet_name, "accepted_qty", "intentional_fixture",
                        "GRN-002 accepted_qty=0 with QUALITY_HOLD is a deliberate 100% hold fixture.",
                        rec.get("accepted_qty"), idx,
                    ))
        if sheet_name in ("ref_supplier_offers", "ref_supplier_contracts"):
            errs += check_date_order(sheet_name, recs, "effective_from", "effective_to")
        if sheet_name == "production_batches":
            errs += check_date_order(sheet_name, recs, "planned_start", "planned_end")
            for idx, rec in enumerate(recs, start=1):
                planned = _to_float(rec.get("planned_output_qty_mt"))
                actual = _to_float(rec.get("actual_output_qty_mt"))
                if planned and planned > 0 and actual is not None and actual / planned < 0.2:
                    errs.append(_info(
                        sheet_name, "actual_output_qty_mt", "low_yield_fixture",
                        f"{rec.get('batch_id')} yield {actual}/{planned} is an extreme underperformance fixture.",
                        actual, idx,
                    ))
        if sheet_name == "shipments":
            errs += check_date_order(sheet_name, recs, "planned_dispatch_at", "estimated_arrival_at")
        if sheet_name == "scenario_templates":
            errs.append(_info(
                sheet_name, "linked_reference_id", "dq01_mislabelled_column",
                "Workbook column linked_reference_id stores severity values in all 15 rows "
                "(DQ-01). Import maps it to severity and sets linked_reference_id=NULL.",
            ))
            errs.append(_info(
                sheet_name, "enabled", "dq02_missing_column",
                "enabled is absent from the workbook; import defaults enabled=TRUE (DQ-02).",
            ))

        expected = EXPECTED_SOURCE_ROW_COUNTS.get(sheet_name)
        if expected is not None and len(recs) != expected:
            errs.append(ValidationError(
                sheet=sheet_name, row=0, column="*", value=len(recs),
                rule="expected_row_count", severity="WARNING",
                message=f"Expected {expected} source rows, found {len(recs)}.",
            ))

        run_sheet(sheet_name, errs)

    empty_runtime = []
    for name in EMPTY_RUNTIME_SHEETS:
        n = len(loaded.get(name, []))
        empty_runtime.append({"sheet_name": name, "row_count": n, "header_only": n == 0})
        if n != 0:
            all_errors.append(ValidationError(
                sheet=name, row=0, column="*", value=n,
                rule="empty_runtime_sheet", severity="WARNING",
                message=f"Runtime sheet {name} was expected empty but has {n} rows.",
            ))

    total_mapped = sum(len(loaded.get(s, [])) for s in SHEET_MAPPINGS)
    workbook_total = sum(len(v) for v in loaded.values())
    # loaded only includes mapped + excluded + empty; count all workbook sheets
    all_sheet_total = 0
    for name in wb.sheetnames:
        if name in loaded:
            all_sheet_total += len(loaded[name])
        else:
            _, recs = sheet_to_records(wb[name])
            all_sheet_total += len(recs)

    summary = {"critical": 0, "warning": 0, "info": 0}
    for e in all_errors:
        key = e.severity.lower()
        if key in summary:
            summary[key] += 1

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "mode": "workbook_only",
        "database_queried": False,
        "workbook_path": str(WORKBOOK_PATH),
        "sheets_inspected": len(sheet_results),
        "total_records_inspected": total_mapped,
        "count_reconciliation": {
            "workbook_sheets": len(wb.sheetnames),
            "workbook_data_rows": all_sheet_total,
            "expected_workbook_data_rows": TOTAL_WORKBOOK_ROWS,
            "mapped_sheets": len(SHEET_MAPPINGS),
            "mapped_rows": total_mapped,
            "expected_mapped_rows": MAPPED_IMPORT_ROWS,
            "metadata_excluded_rows": METADATA_EXCLUDED_ROWS,
            "empty_runtime_sheets": empty_runtime,
            "note": (
                "1677 is every data row in all 40 sheets (including README, dictionaries, "
                "and 0-row runtime templates). 1327 is the 31 mapped import sheets only. "
                "350 metadata rows (8+37+305) are excluded. Six runtime sheets contribute 0 rows."
            ),
        },
        "summary": summary,
        "sheets": sheet_results,
        "errors": [],
    }


def main():
    if not WORKBOOK_PATH.exists():
        print(f"ERROR: Workbook not found at {WORKBOOK_PATH}")
        sys.exit(1)

    print(f"Opening workbook: {WORKBOOK_PATH}")
    wb = openpyxl.load_workbook(str(WORKBOOK_PATH), read_only=True, data_only=True)

    print("Running workbook validation rules …")
    report = validate_all(wb)
    wb.close()

    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(report, f, indent=2, default=str)

    recon = report["count_reconciliation"]
    print(f"\nWorkbook validation report written to: {OUTPUT_PATH}")
    print(f"  Mode               : workbook_only (no PostgreSQL query)")
    print(f"  Sheets inspected   : {report['sheets_inspected']}")
    print(f"  Mapped records     : {report['total_records_inspected']}")
    print(f"  Workbook data rows : {recon['workbook_data_rows']}")
    print(f"  CRITICAL findings  : {report['summary']['critical']}")
    print(f"  WARNING  findings  : {report['summary']['warning']}")
    print(f"  INFO     findings  : {report['summary']['info']}")

    for s in report["sheets"]:
        if s["findings"]:
            print(f"\n  [{s['sheet_name']}] {s['row_count']} rows — {len(s['findings'])} finding group(s)")
            for fnd in s["findings"]:
                print(
                    f"    {fnd['severity']:8s} | {fnd['rule']:30s} | "
                    f"col={fnd['column']:30s} | count={fnd['count']}"
                )

    if report["summary"]["critical"] > 0:
        sys.exit(1)


if __name__ == "__main__":
    main()
